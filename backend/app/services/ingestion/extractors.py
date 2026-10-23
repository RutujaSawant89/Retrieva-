"""Format-specific text extraction into structural blocks.

Each extractor preserves what the format offers: page numbers (PDF), heading styles (DOCX),
heading markup (Markdown/HTML), tables (DOCX/Markdown/HTML/CSV). OCR for scanned PDFs is a
pluggable hook (``OcrEngine``); without one, image-only pages are reported in metadata rather
than silently producing an empty document.
"""

from __future__ import annotations

import csv
import io
import logging
import re
from collections.abc import Callable
from typing import Protocol

from app.services.ingestion.cleaning import normalize_text, strip_page_furniture
from app.services.ingestion.types import Block, ExtractedDocument

logger = logging.getLogger(__name__)


class ExtractionError(Exception):
    """The file could not be parsed (corrupt, encrypted, or unsupported content)."""


class OcrEngine(Protocol):
    def ocr_pdf_page(self, data: bytes, page_index: int) -> str: ...


_ocr_engine: OcrEngine | None = None


def register_ocr_engine(engine: OcrEngine | None) -> None:
    global _ocr_engine
    _ocr_engine = engine


# --------------------------------------------------------------------------------------
# Plain text helpers
# --------------------------------------------------------------------------------------

_NUMBERED_HEADING = re.compile(
    r"^(?:(?:section|article|chapter|part)\s+)?(\d+(?:\.\d+){0,4})\.?\s+[A-Z][^\n]{1,100}$", re.I
)
_LIST_ITEM = re.compile(r"^\s*(?:[-*•▪◦]|\(?[a-z0-9]{1,3}[.)])\s+", re.I)


def decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16"):
        try:
            text = data.decode(encoding)
            if encoding == "utf-16" and not data.startswith((b"\xff\xfe", b"\xfe\xff")):
                continue
            return text
        except UnicodeDecodeError:
            continue
    return data.decode("cp1252", errors="replace")


def _heading_level_for_line(line: str) -> int | None:
    """Heuristic heading detection for unstructured text (PDF/TXT)."""
    stripped = line.strip()
    if (
        not stripped
        or len(stripped) > 110
        or stripped.endswith((".", ",", ";", ":"))
        and not _NUMBERED_HEADING.match(stripped)
    ):
        return None
    match = _NUMBERED_HEADING.match(stripped)
    if match:
        return min(4, match.group(1).count(".") + 1)
    letters = [c for c in stripped if c.isalpha()]
    if len(letters) >= 4 and all(c.isupper() for c in letters) and len(stripped.split()) <= 12:
        return 1
    return None


def blocks_from_plain_text(text: str, page: int | None = None, detect_headings: bool = True) -> list[Block]:
    """Split unstructured text into heading / paragraph / list blocks.

    Paragraph boundaries are blank lines, heading lines, or a short line that ends a sentence
    (typical of PDF text where layout line breaks are not paragraph breaks).
    """
    blocks: list[Block] = []
    lines = text.split("\n")
    widths = sorted(len(ln) for ln in lines if ln.strip())
    typical = widths[int(len(widths) * 0.8)] if widths else 80
    current: list[str] = []

    def flush() -> None:
        if not current:
            return
        is_list = sum(bool(_LIST_ITEM.match(ln)) for ln in current) >= max(1, (len(current) + 1) // 2)
        kind = "list" if is_list else "paragraph"
        joiner = "\n" if is_list else " "
        blocks.append(Block(text=joiner.join(ln.strip() for ln in current), kind=kind, page=page))
        current.clear()

    for line in lines:
        stripped = line.strip()
        if not stripped:
            flush()
            continue
        if detect_headings and (level := _heading_level_for_line(stripped)):
            flush()
            blocks.append(Block(text=stripped, kind="heading", level=level, page=page))
            continue
        if _LIST_ITEM.match(stripped) and current and not _LIST_ITEM.match(current[0]):
            flush()
        current.append(stripped)
        if stripped.endswith((".", "!", "?", ":")) and len(stripped) < 0.7 * typical:
            flush()
    flush()
    return blocks


# --------------------------------------------------------------------------------------
# Format extractors
# --------------------------------------------------------------------------------------


def extract_pdf(data: bytes) -> ExtractedDocument:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception as exc:
                raise ExtractionError("PDF is password protected") from exc
        pages = list(reader.pages)
    except PdfReadError as exc:
        raise ExtractionError(f"Could not read PDF: {exc}") from exc

    page_texts: list[str] = []
    image_only_pages: list[int] = []
    for index, page in enumerate(pages):
        try:
            text = page.extract_text() or ""
        except Exception:  # pypdf can fail on individual malformed pages
            logger.warning("pdf_page_extract_failed", extra={"page": index + 1})
            text = ""
        if not text.strip():
            if _ocr_engine is not None:
                text = _ocr_engine.ocr_pdf_page(data, index)
            else:
                image_only_pages.append(index + 1)
        page_texts.append(normalize_text(text))

    blocks: list[Block] = []
    for index, text in enumerate(strip_page_furniture(page_texts)):
        if text.strip():
            blocks.extend(blocks_from_plain_text(text, page=index + 1))

    meta: dict[str, object] = {}
    try:
        info = reader.metadata or {}
        for key, name in (("/Title", "title"), ("/Author", "author"), ("/Subject", "subject")):
            value = info.get(key)
            if value and str(value).strip():
                meta[name] = str(value).strip()[:300]
    except Exception:
        logger.debug("pdf_metadata_unreadable")
    if image_only_pages:
        meta["pages_without_text"] = image_only_pages
    return ExtractedDocument(blocks=blocks, page_count=len(pages), metadata=meta)


def _table_to_markdown(rows: list[list[str]]) -> str:
    rows = [
        [cell.replace("|", "\\|").replace("\n", " ").strip() for cell in row]
        for row in rows
        if any(c.strip() for c in row)
    ]
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    header, *body = rows
    lines = ["| " + " | ".join(header) + " |", "| " + " | ".join("---" for _ in header) + " |"]
    lines += ["| " + " | ".join(r) + " |" for r in body]
    return "\n".join(lines)


def extract_docx(data: bytes) -> ExtractedDocument:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise ExtractionError(f"Could not read DOCX: {exc}") from exc

    blocks: list[Block] = []
    body = document.element.body
    for child in body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = Paragraph(child, document)
            text = para.text.strip()
            if not text:
                continue
            style = (para.style.name if para.style is not None else "") or ""
            if style.lower().startswith("heading") or style.lower() == "title":
                digits = re.findall(r"\d+", style)
                level = int(digits[0]) if digits else 1
                blocks.append(Block(text=text, kind="heading", level=level))
            elif "list" in style.lower():
                blocks.append(Block(text=f"- {text}", kind="list"))
            else:
                blocks.append(Block(text=text, kind="paragraph"))
        elif tag == "tbl":
            table = Table(child, document)
            rows = [[cell.text for cell in row.cells] for row in table.rows]
            markdown = _table_to_markdown(rows)
            if markdown:
                blocks.append(Block(text=markdown, kind="table"))

    props = document.core_properties
    meta: dict[str, object] = {k: v for k, v in {"title": props.title, "author": props.author}.items() if v}
    return ExtractedDocument(blocks=_merge_list_items(blocks), page_count=None, metadata=meta)


def _merge_list_items(blocks: list[Block]) -> list[Block]:
    merged: list[Block] = []
    for block in blocks:
        if block.kind == "list" and merged and merged[-1].kind == "list":
            merged[-1].text += "\n" + block.text
        else:
            merged.append(block)
    return merged


_MD_HEADING = re.compile(r"^(#{1,6})\s+(.+?)\s*#*\s*$")
_MD_SETEXT = re.compile(r"^(=+|-+)\s*$")


def extract_markdown(data: bytes) -> ExtractedDocument:
    text = decode_text(data)
    blocks: list[Block] = []
    lines = text.splitlines()
    buffer: list[str] = []
    in_code = False
    code: list[str] = []
    meta: dict[str, object] = {}

    def flush() -> None:
        if buffer:
            chunk = "\n".join(buffer).strip()
            if chunk:
                is_table = all(ln.strip().startswith("|") for ln in buffer if ln.strip())
                is_list = sum(bool(_LIST_ITEM.match(ln)) for ln in buffer) >= max(1, len(buffer) // 2)
                kind = "table" if is_table else "list" if is_list else "paragraph"
                blocks.append(
                    Block(text=chunk if kind != "paragraph" else " ".join(chunk.split()), kind=kind)
                )
            buffer.clear()

    for i, line in enumerate(lines):
        if line.strip().startswith("```"):
            if in_code:
                blocks.append(Block(text="\n".join(code), kind="code"))
                code, in_code = [], False
            else:
                flush()
                in_code = True
            continue
        if in_code:
            code.append(line)
            continue
        heading = _MD_HEADING.match(line)
        if heading:
            flush()
            blocks.append(Block(text=heading.group(2).strip(), kind="heading", level=len(heading.group(1))))
            if len(heading.group(1)) == 1 and "title" not in meta:
                meta["title"] = heading.group(2).strip()[:300]
            continue
        if _MD_SETEXT.match(line) and buffer and len(buffer) == 1 and i > 0:
            title = buffer.pop()
            flush()
            blocks.append(Block(text=title.strip(), kind="heading", level=1 if line.strip()[0] == "=" else 2))
            continue
        if not line.strip():
            flush()
            continue
        buffer.append(line)
    if in_code and code:
        blocks.append(Block(text="\n".join(code), kind="code"))
    flush()
    return ExtractedDocument(blocks=blocks, metadata=meta)


def extract_text(data: bytes) -> ExtractedDocument:
    return ExtractedDocument(blocks=blocks_from_plain_text(decode_text(data)))


CSV_ROWS_PER_BLOCK = 25


def extract_csv(data: bytes) -> ExtractedDocument:
    text = decode_text(data)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = list(csv.reader(io.StringIO(text), dialect))
    rows = [r for r in rows if any(cell.strip() for cell in r)]
    if not rows:
        return ExtractedDocument(blocks=[], metadata={"rows": 0})
    header, body = rows[0], rows[1:]
    blocks: list[Block] = []
    # Each record is rendered as "column: value" pairs so it stays meaningful when chunked alone.
    for start in range(0, len(body), CSV_ROWS_PER_BLOCK):
        group = body[start : start + CSV_ROWS_PER_BLOCK]
        lines = []
        for offset, row in enumerate(group):
            pairs = "; ".join(
                f"{h.strip() or f'col{j + 1}'}: {v.strip()}"
                for j, (h, v) in enumerate(zip(header, row, strict=False))
                if v.strip()
            )
            lines.append(f"Row {start + offset + 1}: {pairs}")
        blocks.append(Block(text="\n".join(lines), kind="row"))
    return ExtractedDocument(
        blocks=blocks, metadata={"rows": len(body), "columns": [h.strip() for h in header]}
    )


_HTML_DROP = ("script", "style", "noscript", "nav", "footer", "header", "form", "svg", "iframe", "button")


def extract_html(data: bytes) -> ExtractedDocument:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(decode_text(data), "html.parser")
    for tag in soup(list(_HTML_DROP)):
        tag.decompose()
    meta: dict[str, object] = {}
    if soup.title and soup.title.get_text(strip=True):
        meta["title"] = soup.title.get_text(strip=True)[:300]

    root = soup.find("main") or soup.find("article") or soup.body or soup
    blocks: list[Block] = []
    for element in root.find_all(
        ["h1", "h2", "h3", "h4", "h5", "h6", "p", "li", "table", "pre", "blockquote"]
    ):
        if element.find_parent(["table", "pre"]) is not None and element.name not in ("table", "pre"):
            continue
        if element.name == "li" and element.find_parent("li") is not None:
            continue
        if element.name == "table":
            rows = [
                [cell.get_text(" ", strip=True) for cell in tr.find_all(["th", "td"])]
                for tr in element.find_all("tr")
            ]
            markdown = _table_to_markdown(rows)
            if markdown:
                blocks.append(Block(text=markdown, kind="table"))
            continue
        if element.name == "pre":
            blocks.append(Block(text=element.get_text().strip("\n"), kind="code"))
            continue
        text = " ".join(element.get_text(" ", strip=True).split())
        if not text:
            continue
        if element.name.startswith("h"):
            blocks.append(Block(text=text, kind="heading", level=int(element.name[1])))
        elif element.name == "li":
            blocks.append(Block(text=f"- {text}", kind="list"))
        else:
            blocks.append(Block(text=text, kind="paragraph"))
    return ExtractedDocument(blocks=_merge_list_items(blocks), metadata=meta)


EXTRACTORS: dict[str, Callable[[bytes], ExtractedDocument]] = {
    "pdf": extract_pdf,
    "docx": extract_docx,
    "txt": extract_text,
    "md": extract_markdown,
    "csv": extract_csv,
    "html": extract_html,
}


def extract(file_type: str, data: bytes) -> ExtractedDocument:
    try:
        extractor = EXTRACTORS[file_type]
    except KeyError as exc:
        raise ExtractionError(f"Unsupported file type: {file_type}") from exc
    return extractor(data)
