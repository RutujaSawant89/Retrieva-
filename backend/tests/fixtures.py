"""Builders for realistic test documents."""

from __future__ import annotations

import io

CONTRACT_SECTIONS = [
    (
        "1 Definitions",
        "In this Agreement, Services means the data processing services described in Schedule A. "
        "Customer Data means all data submitted by the Customer to the Provider.",
    ),
    (
        "2 Fees and Payment",
        "The Customer shall pay the annual subscription fee of 48,000 USD within thirty days of "
        "the invoice date. Late payments accrue interest at one percent per month.",
    ),
    (
        "3 Data Protection",
        "The Provider shall encrypt Customer Data at rest using AES-256 and in transit using TLS 1.2 "
        "or higher. Personal data is processed only in the European Economic Area.",
    ),
    (
        "4 Termination",
        "Either party may terminate this Agreement for convenience upon ninety (90) days prior written "
        "notice. Either party may terminate for material breach if the breach is not cured within thirty (30) days.",
    ),
    (
        "5 Limitation of Liability",
        "Neither party's aggregate liability shall exceed the fees paid in the twelve months "
        "preceding the claim, except for breaches of confidentiality.",
    ),
]


def make_pdf(
    sections: list[tuple[str, str]] = CONTRACT_SECTIONS, title: str = "Master Services Agreement"
) -> bytes:
    from fpdf import FPDF

    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    for i, (heading, body) in enumerate(sections):
        if i % 2 == 0:
            pdf.add_page()
            pdf.set_font("Helvetica", size=8)
            pdf.cell(0, 6, f"{title} - Confidential", new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", style="B", size=13)
        pdf.cell(0, 10, heading, new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", size=11)
        pdf.multi_cell(0, 6, body)
        pdf.ln(4)
    return bytes(pdf.output())


def make_docx(sections: list[tuple[str, str]] = CONTRACT_SECTIONS) -> bytes:
    import docx

    document = docx.Document()
    document.add_heading("Master Services Agreement", level=0)
    for heading, body in sections:
        document.add_heading(heading, level=1)
        document.add_paragraph(body)
    table = document.add_table(rows=3, cols=2)
    for r, (a, b) in enumerate([("Tier", "Fee"), ("Standard", "48,000 USD"), ("Premium", "96,000 USD")]):
        table.rows[r].cells[0].text = a
        table.rows[r].cells[1].text = b
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_markdown(sections: list[tuple[str, str]] = CONTRACT_SECTIONS) -> bytes:
    parts = ["# Master Services Agreement", ""]
    for heading, body in sections:
        parts += [f"## {heading}", "", body, ""]
    return "\n".join(parts).encode()
